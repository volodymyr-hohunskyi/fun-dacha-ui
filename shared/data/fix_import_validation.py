#!/usr/bin/env python3
"""
Fix Export/Import validation errors in products_import.xlsx:

- Remove rows with empty product_id from AdditionalImages, ProductSEOKeywords
- Remove rows with empty category_id from Categories, CategorySEOKeywords
- Set store_id=0 when empty in ProductSEOKeywords, CategorySEOKeywords

Usage:
    python3 shared/data/fix_import_validation.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

WORKBOOK_PATH = Path(__file__).resolve().with_name("products_import.xlsx")


def _safe_int(val) -> int | None:
    if val is None or val == "":
        return None
    try:
        return int(val)
    except (TypeError, ValueError):
        return None


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise SystemExit(f"Workbook not found: {WORKBOOK_PATH}")

    wb = load_workbook(WORKBOOK_PATH)

    changes = []

    # AdditionalImages: remove rows with empty product_id
    if "AdditionalImages" in wb.sheetnames:
        ws = wb["AdditionalImages"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        keep = [r for r in rows if _safe_int(r[0]) is not None]
        removed = len(rows) - len(keep)
        if removed > 0:
            ws.delete_rows(2, len(rows))
            for r in keep:
                ws.append(r)
            changes.append(f"AdditionalImages: removed {removed} rows with empty product_id")

    # ProductSEOKeywords: remove rows with empty product_id, fix empty store_id
    if "ProductSEOKeywords" in wb.sheetnames:
        ws = wb["ProductSEOKeywords"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        fixed = []
        for r in rows:
            pid = _safe_int(r[0])
            if pid is None:
                continue
            row = list(r)
            if _safe_int(row[1]) is None:
                row[1] = 0
            fixed.append(row)
        removed = len(rows) - len(fixed)
        if removed > 0 or any(_safe_int(r[1]) is None for r in rows):
            ws.delete_rows(2, len(rows))
            for r in fixed:
                ws.append(r)
            if removed > 0:
                changes.append(f"ProductSEOKeywords: removed {removed} rows with empty product_id")
            if any(_safe_int(r[1]) is None for r in rows):
                changes.append("ProductSEOKeywords: set empty store_id to 0")

    # Categories: remove rows with empty category_id
    if "Categories" in wb.sheetnames:
        ws = wb["Categories"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        keep = [r for r in rows if _safe_int(r[0]) is not None]
        removed = len(rows) - len(keep)
        if removed > 0:
            ws.delete_rows(2, len(rows))
            for r in keep:
                ws.append(r)
            changes.append(f"Categories: removed {removed} rows with empty category_id")

    # CategorySEOKeywords: remove rows with empty category_id, fix empty store_id
    if "CategorySEOKeywords" in wb.sheetnames:
        ws = wb["CategorySEOKeywords"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        fixed = []
        for r in rows:
            cid = _safe_int(r[0])
            if cid is None:
                continue
            row = list(r)
            if _safe_int(row[1]) is None:
                row[1] = 0
            fixed.append(row)
        removed = len(rows) - len(fixed)
        if removed > 0 or any(_safe_int(r[1]) is None for r in rows):
            ws.delete_rows(2, len(rows))
            for r in fixed:
                ws.append(r)
            if removed > 0:
                changes.append(f"CategorySEOKeywords: removed {removed} rows with empty category_id")
            if any(_safe_int(r[1]) is None for r in rows):
                changes.append("CategorySEOKeywords: set empty store_id to 0")

    wb.save(WORKBOOK_PATH)
    for c in changes:
        print(c)
    if not changes:
        print("No validation fixes needed.")


if __name__ == "__main__":
    main()
