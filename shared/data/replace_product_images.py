#!/usr/bin/env python3
"""
Replace product images from shared/images/products to image/catalog/products.
Reads products_import.xlsx Products sheet for model->image mapping.
Source: shared/images/products/{model}.jpg (try model first) or {basename}
Dest: image/catalog/products/{basename}
"""

from pathlib import Path

SHARED_PRODUCTS = Path(__file__).resolve().parent.parent / "images" / "products"
CATALOG_PRODUCTS = Path(__file__).resolve().parent.parent.parent / "image" / "catalog" / "products"
PRODUCTS_IMPORT = Path(__file__).resolve().parent / "products_import.xlsx"

# User's list: model_id from their "ID" column
USER_MODELS = [
    "p401840", "p401846", "p401847", "p401849", "p401850", "p401852", "p401854", "p401855",
    "p401857", "p401859", "p401862", "p401863", "p401864", "p401865", "p401866", "p401867",
    "p401870", "p401874", "p401875", "p401879", "p401880", "p401883", "p401884", "p401885",
    "p401887", "p401890", "p401891", "p401893", "p401896", "p401898", "p401900", "p401902",
    "p401903", "p401904", "p401905", "p401906", "p401907", "p401909", "p401912", "p401914",
    "p401916", "p401917", "p401922", "p401924", "p401925", "p401927", "p401928", "p401929",
    "p401930", "p401932", "p401933", "p401934", "p401939", "p401942", "p401943", "p401944",
    "p401946", "p401947",
    "p402948", "p402949", "p402950", "p402951", "p402952", "p402954", "p402955", "p402956",
    ("p401847", "catalog/products/p400847.jpg"),
    ("p401849", "catalog/products/p400849.jpg"),
    ("p401850", "catalog/products/p400850.jpg"),
    ("p401852", "catalog/products/p400852.jpg"),
    ("p401854", "catalog/products/p400854.jpg"),
    ("p401855", "catalog/products/p400855.jpg"),
    ("p401857", "catalog/products/p400857.jpg"),
    ("p401859", "catalog/products/p400859.jpg"),
    ("p401862", "catalog/products/p400862.jpg"),
    ("p401863", "catalog/products/p400863.jpg"),
    ("p401864", "catalog/products/p400864.jpg"),
    ("p401865", "catalog/products/p400865.jpg"),
    ("p401866", "catalog/products/p400866.jpg"),
    ("p401867", "catalog/products/p400867.jpg"),
    ("p401870", "catalog/products/p400870.jpg"),
    ("p401874", "catalog/products/p400874.jpg"),
    ("p401875", "catalog/products/p400875.jpg"),
    ("p401879", "catalog/products/p400879.jpg"),
    ("p401880", "catalog/products/p400880.jpg"),
    ("p401883", "catalog/products/p400883.jpg"),
    ("p401884", "catalog/products/p400884.jpg"),
    ("p401885", "catalog/products/p400885.jpg"),
    ("p401887", "catalog/products/p400887.jpg"),
    ("p401890", "catalog/products/p400890.jpg"),
    ("p401891", "catalog/products/p400891.jpg"),
    ("p401893", "catalog/products/p400893.jpg"),
    ("p401896", "catalog/products/p400896.jpg"),
    ("p401898", "catalog/products/p400898.jpg"),
    ("p401900", "catalog/products/p400900.jpg"),
    ("p401902", "catalog/products/p400902.jpg"),
    ("p401903", "catalog/products/p400903.jpg"),
    ("p401904", "catalog/products/p400904.jpg"),
    ("p401905", "catalog/products/p400905.jpg"),
    ("p401906", "catalog/products/p400906.jpg"),
    ("p401907", "catalog/products/p400907.jpg"),
    ("p401909", "catalog/products/p400909.jpg"),
    ("p401912", "catalog/products/p400912.jpg"),
    ("p401914", "catalog/products/p400914.jpg"),
    ("p401916", "catalog/products/p400916.jpg"),
    ("p401917", "catalog/products/p400917.jpg"),
    ("p401922", "catalog/products/p400922.jpg"),
    ("p401924", "catalog/products/p400924.jpg"),
    ("p401925", "catalog/products/p400925.jpg"),
    ("p401927", "catalog/products/p400927.jpg"),
    ("p401928", "catalog/products/p400928.jpg"),
    ("p401929", "catalog/products/p400929.jpg"),
    ("p401930", "catalog/products/p400930.jpg"),
    ("p401932", "catalog/products/p400932.jpg"),
    ("p401933", "catalog/products/p400933.jpg"),
    ("p401934", "catalog/products/p400934.jpg"),
    ("p401939", "catalog/products/p400939.jpg"),
    ("p401942", "catalog/products/p400942.jpg"),
    ("p401943", "catalog/products/p400943.jpg"),
    ("p401944", "catalog/products/p400944.jpg"),
    ("p401946", "catalog/products/p400946.jpg"),
    ("p401947", "catalog/products/p400947.jpg"),
    ("p402948", "catalog/products/p400948.jpg"),
    ("p402949", "catalog/products/p400949.jpg"),
    ("p402950", "catalog/products/p400950.jpg"),
    ("p402951", "catalog/products/p400951.jpg"),
    ("p402952", "catalog/products/p400952.jpg"),
    ("p402954", "catalog/products/p400954.jpg"),
    ("p402955", "catalog/products/p400955.jpg"),
    ("p402956", "catalog/products/p400956.jpg"),
    ("p402957", "catalog/products/p400957.jpg"),
    ("p402958", "catalog/products/p400958.jpg"),
    ("p402962", "catalog/products/p400962.jpg"),
    ("p402963", "catalog/products/p400963.jpg"),
    ("p402966", "catalog/products/p400966.jpg"),
    ("p402969", "catalog/products/p400969.jpg"),
    ("p402970", "catalog/products/p400970.jpg"),
    ("p402971", "catalog/products/p400971.jpg"),
    ("p402973", "catalog/products/p400973.jpg"),
    ("p402975", "catalog/products/p400975.jpg"),
    ("p402980", "catalog/products/p400980.jpg"),
    ("p402981", "catalog/products/p400981.jpg"),
    ("p402982", "catalog/products/p400982.jpg"),
    ("p402984", "catalog/products/p400984.jpg"),
    ("p402985", "catalog/products/p400985.jpg"),
    ("p402986", "catalog/products/p400986.jpg"),
    ("p402987", "catalog/products/p400987.jpg"),
    ("p402990", "catalog/products/p400990.jpg"),
    ("p402991", "catalog/products/p400991.jpg"),
    ("p402992", "catalog/products/p400992.jpg"),
    ("p402994", "catalog/products/p400994.jpg"),
    ("p402996", "catalog/products/p400996.jpg"),
    ("p4021005", "catalog/products/p4001005.jpg"),
    ("p4021010", "catalog/products/p4001010.jpg"),
    ("p4021011", "catalog/products/p4001011.jpg"),
    ("p4021014", "catalog/products/p4001014.jpg"),
    ("p4021015", "catalog/products/p4001015.jpg"),
    ("p4021019", "catalog/products/p4001019.jpg"),
    ("p4021021", "catalog/products/p4001021.jpg"),
    ("p4021023", "catalog/products/p4001023.jpg"),
    ("p4021026", "catalog/products/p4001026.jpg"),
    ("p4021028", "catalog/products/p4001028.jpg"),
    ("p4021029", "catalog/products/p4001029.jpg"),
    ("p4021034", "catalog/products/p4001034.jpg"),
    ("p4021038", "catalog/products/p4001038.jpg"),
    ("p4021040", "catalog/products/p4001040.jpg"),
    ("p4021041", "catalog/products/p4001041.jpg"),
    ("p4021045", "catalog/products/p4001045.jpg"),
    ("p4021046", "catalog/products/p4001046.jpg"),
    ("p4021058", "catalog/products/p4001058.jpg"),
    ("p4021060", "catalog/products/p4001060.jpg"),
    ("p4021061", "catalog/products/p4001061.jpg"),
    ("p4021066", "catalog/products/p4001066.jpg"),
    ("p4021071", "catalog/products/p4001071.jpg"),
    ("p4021073", "catalog/products/p4001073.jpg"),
    ("p4021074", "catalog/products/p4001074.jpg"),
    ("p4021075", "catalog/products/p4001075.jpg"),
    ("p4021077", "catalog/products/p4001077.jpg"),
    ("p4021082", "catalog/products/p4001082.jpg"),
    ("p4021090", "catalog/products/p4001090.jpg"),
    ("p4021094", "catalog/products/p4001094.jpg"),
    ("p4021096", "catalog/products/p4001096.jpg"),
    ("p4021098", "catalog/products/p4001098.jpg"),
    ("p4021099", "catalog/products/p4001099.jpg"),
    ("p4021102", "catalog/products/p4001102.jpg"),
    ("p4021103", "catalog/products/p4001103.jpg"),
    ("p4021105", "catalog/products/p4001105.jpg"),
    ("p4021106", "catalog/products/p4001106.jpg"),
    ("p4021107", "catalog/products/p4001107.jpg"),
    ("p4021108", "catalog/products/p4001108.jpg"),
    ("p4021109", "catalog/products/p4001109.jpg"),
    ("p4021112", "catalog/products/p4001112.jpg"),
    ("p4021123", "catalog/products/p4001123.jpg"),
    ("p4021124", "catalog/products/p4001124.jpg"),
    "p4021127",
]


def load_products_mapping():
    """Load model -> image_name from products_import.xlsx."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {}
    if not PRODUCTS_IMPORT.exists():
        return {}
    wb = load_workbook(PRODUCTS_IMPORT, read_only=True)
    if "Products" not in wb.sheetnames:
        return {}
    ws = wb["Products"]
    headers = [c.value for c in ws[1]]
    model_idx = headers.index("model") if "model" in headers else 7
    img_idx = headers.index("image_name") if "image_name" in headers else 9
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        m = row[model_idx]
        img = row[img_idx]
        if m and img:
            out[str(m).strip()] = str(img).strip()
    return out


def main():
    import shutil

    mapping = load_products_mapping()
    # Normalize USER_MODELS (handle tuple leftovers, dedupe)
    models = list(dict.fromkeys(m[0] if isinstance(m, tuple) else m for m in USER_MODELS))
    # Build items: for each user model, get catalog path from products_import or infer p400xxx
    items = []
    for model in models:
        catalog_path = mapping.get(model)
        if not catalog_path:
            # Infer: p401840 -> catalog/products/p400840.jpg
            s = model.replace("p402", "p400").replace("p401", "p400")
            catalog_path = f"catalog/products/{s}.jpg"
        else:
            catalog_path = catalog_path.replace("\\", "/")
        items.append((model, catalog_path))

    CATALOG_PRODUCTS.mkdir(parents=True, exist_ok=True)
    replaced = []
    not_found = []

    for model_id, catalog_path in items:
        basename = Path(catalog_path).name
        # Try source as model name first, then as catalog basename
        src = SHARED_PRODUCTS / f"{model_id}.jpg"
        if not src.exists():
            src = SHARED_PRODUCTS / basename
        dest = CATALOG_PRODUCTS / basename

        if src.exists():
            shutil.copy2(src, dest)
            replaced.append((model_id, basename))
        else:
            not_found.append((model_id, basename))

    print("=== REPLACED ===")
    for m, b in replaced:
        print(f"  {m} -> {b}")
    print(f"\nTotal replaced: {len(replaced)}")

    print("\n=== NOT FOUND in shared/images/products ===")
    for m, b in not_found:
        print(f"  {m} (need {m}.jpg for catalog {b})")
    print(f"\nTotal not replaced: {len(not_found)}")


if __name__ == "__main__":
    main()
