#!/usr/bin/env python3
"""
Fix product image_name in products_import.xlsx:
1. If current image exists in catalog -> keep
2. If not, try catalog/products/{model}.jpg -> if exists, update image_name
3. If shared has {model}.jpg, copy to catalog and update image_name

Run: python3 shared/data/fix_product_images.py
"""

from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent.parent
CATALOG = ROOT / "image" / "catalog" / "products"
SHARED = ROOT / "shared" / "images" / "products"
PRODUCTS_IMPORT = Path(__file__).resolve().parent / "products_import.xlsx"


def main():
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("Need openpyxl: pip install openpyxl")
        return

    wb = load_workbook(PRODUCTS_IMPORT)
    ws = wb["Products"]
    headers = [c.value for c in ws[1]]
    model_idx = headers.index("model")
    img_idx = headers.index("image_name")

    catalog_ok = 0
    fixed_by_model = []
    fixed_by_copy = []
    still_missing = []

    for row_idx in range(2, ws.max_row + 1):
        model = ws.cell(row_idx, model_idx + 1).value
        img = ws.cell(row_idx, img_idx + 1).value
        if not model:
            continue
        model = str(model).strip()
        img = str(img).strip().replace("\\", "/") if img else ""
        basename = img.split("/")[-1] if img and "/" in img else (img or "")

        path_orig = CATALOG / basename if basename else None
        path_model = CATALOG / f"{model}.jpg"
        path_shared = SHARED / f"{model}.jpg"

        if path_orig and path_orig.exists():
            catalog_ok += 1
            continue

        # Try model.jpg (id.jpg) - often exists when image_name has typo
        if path_model.exists():
            new_img = f"catalog/products/{model}.jpg"
            ws.cell(row_idx, img_idx + 1).value = new_img
            fixed_by_model.append((model, basename or "?", new_img))
            continue

        if path_shared.exists():
            import shutil
            path_model.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path_shared, path_model)
            new_img = f"catalog/products/{model}.jpg"
            ws.cell(row_idx, img_idx + 1).value = new_img
            fixed_by_copy.append((model, basename or "?", new_img))
            continue

        still_missing.append((model, basename or "?"))

    wb.save(PRODUCTS_IMPORT)

    print("=== Already OK (image exists in catalog) ===")
    print(f"  {catalog_ok} products")

    print("\n=== Fixed: image_name -> model.jpg (file existed) ===")
    print(f"  {len(fixed_by_model)} products")
    for m, old, new in fixed_by_model[:20]:
        print(f"    {m}: {old} -> {new}")
    if len(fixed_by_model) > 20:
        print(f"    ... and {len(fixed_by_model) - 20} more")

    print("\n=== Fixed: copied from shared to catalog ===")
    print(f"  {len(fixed_by_copy)} products")
    for m, old, new in fixed_by_copy[:20]:
        print(f"    {m}: copied -> {new}")
    if len(fixed_by_copy) > 20:
        print(f"    ... and {len(fixed_by_copy) - 20} more")

    print("\n=== Still missing (no catalog nor shared image) ===")
    print(f"  {len(still_missing)} products")
    for m, old in still_missing[:30]:
        print(f"    {m} (was {old})")
    if len(still_missing) > 30:
        print(f"    ... and {len(still_missing) - 30} more")


if __name__ == "__main__":
    main()
