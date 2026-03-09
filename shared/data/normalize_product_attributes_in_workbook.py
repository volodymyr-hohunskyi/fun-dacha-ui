#!/usr/bin/env python3
"""
Normalize ProductAttributes sheet in products_import.xlsx: split combined values
into separate rows. Each row has only one attribute value.

Input:  product_id=1, Колір, text="Червоний, Жовтий"
Output: two rows: (1, Колір, Червоний) and (1, Колір, Жовтий)

Also adds ProductFilters, FilterGroups, Filters for OpenCart filtering.

Usage:
    python3 shared/data/normalize_product_attributes_in_workbook.py
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook

try:
    from attribute_schema import (
        ATTRIBUTE_DEFINITIONS,
        ATTRIBUTE_LOOKUP,
        CHARACTERISTICS_GROUP,
        FILTER_VALUE_TAXONOMY,
        normalize_filter_value,
    )
except ImportError:
    from .attribute_schema import (
        ATTRIBUTE_DEFINITIONS,
        ATTRIBUTE_LOOKUP,
        CHARACTERISTICS_GROUP,
        FILTER_VALUE_TAXONOMY,
        normalize_filter_value,
    )

WORKBOOK_PATH = Path(__file__).resolve().with_name("products_import.xlsx")
SHEET_NAME = "ProductAttributes"
TARGET_LANGUAGE = "uk-ua"


def split_values(raw: str | None) -> List[str]:
    if not raw:
        return []
    parts = [p.strip() for p in str(raw).split(",")]
    return [p for p in parts if p]


def pad(values: List[str], target: int, fallback: str = "") -> List[str]:
    if not values:
        return [fallback] * target if target else []
    if len(values) >= target:
        return values[:target]
    return values + [fallback] * (target - len(values))


def resolve_attr_name(attribute: str) -> str:
    """Resolve attribute to canonical uk-ua name."""
    key = (attribute or "").strip().lower()
    if key:
        defn = ATTRIBUTE_LOOKUP.get(key)
        if defn:
            return defn["names"].get(TARGET_LANGUAGE, defn["names"]["en-gb"])
    return (attribute or "").strip()


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise SystemExit(f"Workbook not found: {WORKBOOK_PATH}")

    wb = load_workbook(WORKBOOK_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Sheet '{SHEET_NAME}' not found")

    ws = wb[SHEET_NAME]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    # One row per value: (product_id, attr_ua, en, ua, ru)
    pa_rows: List[Tuple[int, str, str, str, str]] = []
    pf_rows: List[Tuple[int, str, str]] = []
    taxonomy_seen: Dict[str, List[str]] = {}

    for row in rows:
        if len(row) < 6:
            continue
        product_id_raw, group, attribute, text_en, text_ua, text_ru = row[:6]
        try:
            product_id = int(product_id_raw)
        except (TypeError, ValueError):
            continue

        attr_ua = resolve_attr_name(attribute or group)
        if not attr_ua:
            continue

        en_tokens = split_values(text_en) or ([str(text_en or "").strip()] if text_en else [])
        ua_tokens = split_values(text_ua) or ([str(text_ua or "").strip()] if text_ua else [])
        ru_tokens = split_values(text_ru) or ([str(text_ru or "").strip()] if text_ru else [])

        ua_tokens = pad(ua_tokens, len(en_tokens), (text_ua or "").strip() if ua_tokens else "")
        ru_tokens = pad(ru_tokens, len(en_tokens), (text_ru or "").strip() if ru_tokens else "")

        for i, en_val in enumerate(en_tokens):
            ua_val = ua_tokens[i] if i < len(ua_tokens) else en_val
            ru_val = ru_tokens[i] if i < len(ru_tokens) else ua_val
            if not en_val and not ua_val:
                continue
            en_val = en_val or ua_val
            ua_val = ua_val or en_val
            ru_val = ru_val or ua_val
            pa_rows.append((product_id, attr_ua, en_val, ua_val, ru_val))
            canon_ua = normalize_filter_value(attr_ua, ua_val) if ua_val else None
            if canon_ua:
                pf_rows.append((product_id, attr_ua, canon_ua))
                taxonomy_seen.setdefault(attr_ua, [])
                if canon_ua not in taxonomy_seen[attr_ua]:
                    taxonomy_seen[attr_ua].append(canon_ua)

    # Rewrite ProductAttributes: one row per value
    ws.delete_rows(2, max(1, ws.max_row - 1))
    attr_group = CHARACTERISTICS_GROUP["names"][TARGET_LANGUAGE]
    for product_id, attr_ua, en_val, ua_val, ru_val in sorted(pa_rows, key=lambda x: (x[0], x[1], x[3])):
        ws.append((product_id, attr_group, attr_ua, en_val, ua_val, ru_val))

    # Add FilterGroups and Filters first (needed for filter_id mapping)
    all_values: Dict[str, List[str]] = dict(FILTER_VALUE_TAXONOMY)
    for g, vals in taxonomy_seen.items():
        for v in vals:
            if v not in all_values.setdefault(g, []):
                all_values[g].append(v)

    group_to_id: Dict[str, int] = {}
    group_sort: Dict[str, int] = {}
    for idx, defn in enumerate(ATTRIBUTE_DEFINITIONS):
        name_ua = defn["names"].get(TARGET_LANGUAGE, defn["names"]["en-gb"])
        if name_ua in all_values:
            group_to_id[name_ua] = len(group_to_id) + 1
            group_sort[name_ua] = defn.get("sort_order", idx)
    for name_ua in taxonomy_seen:
        if name_ua not in group_to_id:
            group_to_id[name_ua] = len(group_to_id) + 1
            group_sort[name_ua] = 99

    if group_to_id:
        if "FilterGroups" not in wb.sheetnames:
            ws_fg = wb.create_sheet("FilterGroups")
            ws_fg.append(["filter_group_id", "sort_order", "name(en-gb)", "name(uk-ua)", "name(ru-ru)"])
        else:
            ws_fg = wb["FilterGroups"]
        # Clear old data, keep header
        while ws_fg.max_row > 1:
            ws_fg.delete_rows(2, 1)
        for name_ua in sorted(group_to_id.keys(), key=lambda x: (group_sort.get(x, 99), x)):
            gid = group_to_id[name_ua]
            defn = next((d for d in ATTRIBUTE_DEFINITIONS if d["names"].get(TARGET_LANGUAGE) == name_ua), None)
            en = defn["names"]["en-gb"] if defn else name_ua
            ru = defn["names"].get("ru-ru", name_ua) if defn else name_ua
            ws_fg.append((gid, group_sort.get(name_ua, 99), en, name_ua, ru))

        if "Filters" not in wb.sheetnames:
            ws_f = wb.create_sheet("Filters")
            ws_f.append(["filter_id", "filter_group_id", "sort_order", "name(en-gb)", "name(uk-ua)", "name(ru-ru)"])
        else:
            ws_f = wb["Filters"]
        while ws_f.max_row > 1:
            ws_f.delete_rows(2, 1)
        fid = 1
        name_to_ids: Dict[Tuple[str, str], Tuple[int, int]] = {}
        for name_ua in sorted(group_to_id.keys(), key=lambda x: (group_sort.get(x, 99), x)):
            gid = group_to_id[name_ua]
            for si, val in enumerate(all_values.get(name_ua, []), 1):
                ws_f.append((fid, gid, si, val, val, val))
                name_to_ids[(name_ua, val)] = (gid, fid)
                fid += 1

        # ProductFilters: product_id, filter_group_id, filter_id (OpenCart-compatible header)
        if "ProductFilters" in wb.sheetnames:
            ws_pf = wb["ProductFilters"]
            ws_pf.delete_rows(1, ws_pf.max_row)
        else:
            ws_pf = wb.create_sheet("ProductFilters", wb.sheetnames.index(SHEET_NAME) + 1)
        ws_pf.append(["product_id", "filter_group_id", "filter_id"])
        for product_id, fg_name, f_name in pf_rows:
            ids = name_to_ids.get((fg_name, f_name))
            if ids:
                gid, fid = ids
                ws_pf.append((product_id, gid, fid))
    else:
        # No FilterGroups – use names (product_id, filter_group, filter)
        if "ProductFilters" in wb.sheetnames:
            ws_pf = wb["ProductFilters"]
            ws_pf.delete_rows(1, ws_pf.max_row)
        else:
            ws_pf = wb.create_sheet("ProductFilters", wb.sheetnames.index(SHEET_NAME) + 1)
        ws_pf.append(["product_id", "filter_group", "filter"])
        for r in pf_rows:
            ws_pf.append(r)

    wb.save(WORKBOOK_PATH)
    print(f"ProductAttributes: {len(rows)} rows → {len(pa_rows)} rows (one value per row)")
    print(f"ProductFilters: {len(pf_rows)} rows")
    if group_to_id:
        print(f"Filter groups: {list(group_to_id.keys())}")


if __name__ == "__main__":
    main()
