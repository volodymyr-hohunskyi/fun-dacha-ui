#!/usr/bin/env python3
"""
Remove orphan product images from catalog - files not referenced by any product.
Reads products_import.xlsx for model and image_name. Keeps only images that
at least one product needs.
"""

from pathlib import Path
from openpyxl import load_workbook

CATALOG = Path(__file__).resolve().parent.parent.parent / "image" / "catalog" / "products"
PRODUCTS_IMPORT = Path(__file__).resolve().parent / "products_import.xlsx"


def main():
    wb = load_workbook(PRODUCTS_IMPORT, read_only=True)
    ws = wb["Products"]
    h = [c.value for c in ws[1]]
    model_idx = h.index("model")
    img_idx = h.index("image_name")

    expected = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        m, img = row[model_idx], row[img_idx]
        if m:
            expected.add(f"{str(m).strip()}.jpg")
        if img:
            b = str(img).strip().replace("\\", "/").split("/")[-1]
            if b:
                expected.add(b)

    def is_needed(name: str) -> bool:
        if name in expected:
            return True
        if "_" in name:  # e.g. p500490_1.jpg -> base p500490.jpg
            base = name.split("_")[0] + ".jpg"
            return base in expected
        return False

    removed = 0
    for f in list(CATALOG.iterdir()):
        if not f.is_file():
            continue
        if not is_needed(f.name):
            f.unlink()
            removed += 1
            if removed <= 30:
                print(f"  Removed: {f.name}")

    print(f"\nRemoved {removed} orphan images from catalog.")


if __name__ == "__main__":
    import sys
    if "--dry-run" in sys.argv:
        print("Dry run - would remove orphan images. Run without --dry-run to execute.")
        expected = set()
        wb = load_workbook(PRODUCTS_IMPORT, read_only=True)
        ws = wb["Products"]
        h = [c.value for c in ws[1]]
        model_idx = h.index("model")
        img_idx = h.index("image_name")
        for row in ws.iter_rows(min_row=2, values_only=True):
            m, img = row[model_idx], row[img_idx]
            if m:
                expected.add(f"{str(m).strip()}.jpg")
            if img:
                b = str(img).strip().replace("\\", "/").split("/")[-1]
                if b:
                    expected.add(b)

        def need(name):
            return name in expected or ("_" in name and name.split("_")[0] + ".jpg" in expected)
        orphans = [f.name for f in CATALOG.iterdir() if f.is_file() and not need(f.name)]
        print(f"Would remove {len(orphans)} files")
    else:
        main()
