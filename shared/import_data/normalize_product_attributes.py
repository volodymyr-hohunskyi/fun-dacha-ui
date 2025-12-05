#!/usr/bin/env python3
"""
Normalize ProductAttributes so every row belongs to the single attribute group
"Характеристики" and uses localized characteristic names (Колір, Ріст, etc.).

Usage:
    python3 shared/import_data/normalize_product_attributes.py
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook

try:
    from attribute_schema import ATTRIBUTE_LOOKUP, CHARACTERISTICS_GROUP
except ImportError:  # pragma: no cover - fallback when executed as a module
    from .attribute_schema import (  # type: ignore
        ATTRIBUTE_LOOKUP,
        CHARACTERISTICS_GROUP,
    )

WORKBOOK_PATH = Path(__file__).with_name("products_import.xlsx")
SHEET_NAME = "ProductAttributes"
TARGET_LANGUAGE = "uk-ua"  # must match store default language for attributes
ATTRIBUTE_GROUP_NAME = CHARACTERISTICS_GROUP["names"][TARGET_LANGUAGE]


def split_values(raw: str | None) -> List[str]:
    if not raw:
        return []
    parts = [part.strip() for part in str(raw).split(",")]
    return [part for part in parts if part]


def pad(values: List[str], target: int, fallback: str = "") -> List[str]:
    if not values:
        return [fallback] * target
    if len(values) >= target:
        return values[:target]
    return values + [fallback] * (target - len(values))


def resolve_attribute_definition(group: str, attribute: str):
    group_key = (group or "").strip().lower()
    if group_key:
        definition = ATTRIBUTE_LOOKUP.get(group_key)
        if definition:
            return definition
    attr_key = (attribute or "").strip().lower()
    if attr_key:
        definition = ATTRIBUTE_LOOKUP.get(attr_key)
        if definition:
            return definition
    return None


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise SystemExit(f"Workbook not found: {WORKBOOK_PATH}")

    wb = load_workbook(WORKBOOK_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Worksheet '{SHEET_NAME}' missing in workbook")

    ws = wb[SHEET_NAME]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    aggregated: Dict[Tuple[int, str], Dict[str, List[str]]] = {}
    for product_id, group, attribute, text_en, text_ua, text_ru in rows:
        if product_id is None:
            continue
        definition = resolve_attribute_definition(group, attribute)
        if not definition:
            raise SystemExit(
                f"Unknown attribute group '{group}' / attribute '{attribute}' in ProductAttributes worksheet."
            )

        en_tokens = split_values(text_en) or [(text_en or "").strip()]
        ua_tokens = split_values(text_ua)
        ru_tokens = split_values(text_ru)

        ua_tokens = pad(ua_tokens, len(en_tokens), (text_ua or "").strip())
        ru_tokens = pad(ru_tokens, len(en_tokens), (text_ru or "").strip())

        localized_attr = definition["names"].get(
            TARGET_LANGUAGE, definition["names"]["en-gb"]
        )
        key = (int(product_id), localized_attr)
        bucket = aggregated.setdefault(
            key,
            {
                "en": [],
                "ua": [],
                "ru": [],
            },
        )
        for idx, en_value in enumerate(en_tokens):
            if not en_value:
                continue
            bucket["en"].append(en_value)
            bucket["ua"].append(ua_tokens[idx])
            bucket["ru"].append(ru_tokens[idx])

    # clear existing rows (except header)
    ws.delete_rows(2, ws.max_row)

    normalized_rows = 0
    for (product_id, localized_attr), texts in sorted(aggregated.items()):
        en_text = ", ".join(dict.fromkeys(filter(None, texts["en"])))
        ua_text = ", ".join(dict.fromkeys(filter(None, texts["ua"])))
        ru_text = ", ".join(dict.fromkeys(filter(None, texts["ru"])))
        ws.append(
            (
                product_id,
                ATTRIBUTE_GROUP_NAME,
                localized_attr,
                en_text,
                ua_text,
                ru_text,
            )
        )
        normalized_rows += 1

    wb.save(WORKBOOK_PATH)
    print(f"Normalized {len(rows)} source rows into {normalized_rows} attribute rows")


if __name__ == "__main__":
    main()
