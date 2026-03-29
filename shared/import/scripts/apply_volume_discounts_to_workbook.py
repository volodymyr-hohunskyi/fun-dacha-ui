#!/usr/bin/env python3
"""
Add / refresh the **Discounts** worksheet on `products_import.xlsx` for volume pricing.

Rules (OpenCart 4.1+ product_discount, non-special rows):
  - quantity 2, priority 0, 5% off (type P)
  - quantity 5, priority 0, 10% off (type P)

Applies to every product whose **categories** cell (comma-separated category_id list)
intersects TARGET_CATEGORY_IDS.

Usage (from repo root):
  python3 shared/import/scripts/apply_volume_discounts_to_workbook.py

Requires: pip install openpyxl
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Install openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

# Category IDs to include (seed catalog groups)
TARGET_CATEGORY_IDS = frozenset(
    {
        100,
        101,
        102,
        103,
        110,
        120,
        121,
        122,
        123,
        124,
        130,
        140,
        141,
        142,
        150,
        160,
        170,
        180,
        190,
        200,
        210,
        220,
        230,
        240,
        250,
        270,
        280,
        281,
        282,
        290,
        330,
        331,
        332,
    }
)

# (quantity, priority, percent, type) — percent used when type == P
DISCOUNT_TIERS = (
    (2, 0, 5, "P"),
    (5, 0, 10, "P"),
)

# Must match a customer group **name** in OpenCart (Sales → Customers → Customer Groups)
DEFAULT_CUSTOMER_GROUP = "Default"

DISCOUNTS_HEADER = (
    "product_id",
    "customer_group",
    "quantity",
    "priority",
    "price",
    "type",
    "special",
    "date_start",
    "date_end",
)


def parse_categories(raw: object) -> set[int]:
    if raw is None:
        return set()
    s = str(raw).strip()
    if not s:
        return set()
    out: set[int] = set()
    for part in s.replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        # allow "100" or path fragments; take numeric tokens
        for token in part.split("/"):
            token = token.strip()
            if token.isdigit():
                out.add(int(token))
    return out


def product_ids_for_categories(ws, headers: list) -> list[int]:
    hi = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}
    if "product_id" not in hi or "categories" not in hi:
        raise SystemExit("Products sheet must have product_id and categories columns")

    pi, ci = hi["product_id"], hi["categories"]
    ids: list[int] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[pi] is None:
            continue
        try:
            pid = int(row[pi])
        except (TypeError, ValueError):
            continue
        cats = parse_categories(row[ci] if ci < len(row) else None)
        if cats & TARGET_CATEGORY_IDS:
            ids.append(pid)

    ids.sort()
    return ids


def build_discount_rows(product_ids: list[int]) -> list[tuple]:
    rows = [DISCOUNTS_HEADER]
    for pid in product_ids:
        for qty, prio, pct, typ in DISCOUNT_TIERS:
            rows.append(
                (
                    pid,
                    DEFAULT_CUSTOMER_GROUP,
                    qty,
                    prio,
                    pct,
                    typ,
                    "false",
                    "0000-00-00",
                    "0000-00-00",
                )
            )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "workbook",
        nargs="?",
        default=str(Path(__file__).resolve().parents[1] / "products_import.xlsx"),
        help="Path to products_import.xlsx",
    )
    args = parser.parse_args()
    path = Path(args.workbook)

    if not path.is_file():
        raise SystemExit(f"Missing workbook: {path}")

    wb = openpyxl.load_workbook(path)
    if "Products" not in wb.sheetnames:
        raise SystemExit("Workbook has no Products sheet")

    ws = wb["Products"]
    headers = [c.value for c in ws[1]]
    product_ids = product_ids_for_categories(ws, headers)

    if "Discounts" in wb.sheetnames:
        wb.remove(wb["Discounts"])

    idx = wb.sheetnames.index("Products") + 1
    dws = wb.create_sheet("Discounts", idx)

    for r, row in enumerate(build_discount_rows(product_ids), start=1):
        for c, val in enumerate(row, start=1):
            dws.cell(row=r, column=c, value=val)

    wb.save(path)
    print(
        f"Wrote {path.name}: Discounts sheet with {len(product_ids)} products × {len(DISCOUNT_TIERS)} tiers = {len(product_ids) * len(DISCOUNT_TIERS)} rows"
    )


if __name__ == "__main__":
    main()
