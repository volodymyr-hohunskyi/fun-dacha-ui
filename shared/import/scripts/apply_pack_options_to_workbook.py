#!/usr/bin/env python3
"""
Build **pack-size product options** (1 / 2 / 5) in `products_import.xlsx` for Export/Import.

Adds / refreshes:
  - **Options** + **OptionValues** (global definitions)
  - **ProductOptions** + **ProductOptionValues** per matching product

The **`Discounts`** sheet is **kept**: legacy **quantity > 1** volume rows are removed so they do not stack with pack options; **quantity = 1** rows (e.g. **special** sale price for OpenCart 4.1+) stay — align those with your **SalesCalendar** / pricing process.

Pricing: **Products.price** is treated as the **per-unit selling price** (after sale in your sheet).
Option price modifiers stack so **cart line total** = base_unit + option_delta = N × unit × (1 − tier%).

  - 1 pack: +0
  - 2 packs: −5% per unit  →  delta = 2×0.95×unit − unit
  - 5 packs: −10% per unit → delta = 5×0.90×unit − unit

Same **category filter** as the old volume-discount script.

Requires: pip install openpyxl

Usage:
  python3 shared/import/scripts/apply_pack_options_to_workbook.py [path/to/products_import.xlsx]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Install openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

TARGET_CATEGORY_IDS = frozenset(
    {
        100,
        101,
        102,
        103,
        110,
        120,
        121,
        122,
        123,
        124,
        130,
        140,
        141,
        142,
        150,
        160,
        170,
        180,
        190,
        200,
        210,
        220,
        230,
        240,
        250,
        270,
        280,
        281,
        282,
        290,
        330,
        331,
        332,
    }
)

# Stable IDs — avoid collisions with existing options in DB when using numeric IDs in sheet
OPTION_ID = 92001
OPTION_VALUE_IDS = (920011, 920012, 920013)

# Tier: (packs, percent_off_per_unit_vs_base)
TIERS = (
    (1, 0.0),
    (2, 5.0),
    (5, 10.0),
)

# Must match admin **default language** `name(...)` column for ProductOptions column "option"
# Must match OpenCart **config_language** (admin default). Used for ProductOptions / ProductOptionValues name columns.
DEFAULT_LANG_CODE = "en-gb"

LANGS = ("en-gb", "uk-ua", "ru-ru")
OPTION_NAMES = {
    "en-gb": "Pack quantity",
    "uk-ua": "Кількість упаковок",
    "ru-ru": "Количество упаковок",
}
VALUE_NAMES = (
    {
        "en-gb": "1 pack",
        "uk-ua": "1 уп",
        "ru-ru": "1 уп.",
    },
    {
        "en-gb": "2 packs",
        "uk-ua": "2 уп",
        "ru-ru": "2 уп.",
    },
    {
        "en-gb": "5 packs",
        "uk-ua": "5 уп",
        "ru-ru": "5 уп.",
    },
)


def parse_categories(raw: object) -> set[int]:
    if raw is None:
        return set()
    s = str(raw).strip()
    if not s:
        return set()
    out: set[int] = set()
    for part in s.replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        for token in part.split("/"):
            token = token.strip()
            if token.isdigit():
                out.add(int(token))
    return out


def product_ids_for_categories(ws, headers: list) -> list[int]:
    hi = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}
    if "product_id" not in hi or "categories" not in hi:
        raise SystemExit("Products sheet must have product_id and categories columns")
    pi, ci = hi["product_id"], hi["categories"]
    price_i = hi.get("price")
    ids: list[int] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[pi] is None:
            continue
        try:
            pid = int(row[pi])
        except (TypeError, ValueError):
            continue
        cats = parse_categories(row[ci] if ci < len(row) else None)
        if not (cats & TARGET_CATEGORY_IDS):
            continue
        if price_i is not None and price_i < len(row):
            try:
                p = float(row[price_i] or 0)
            except (TypeError, ValueError):
                p = 0.0
            if p <= 0:
                continue
        ids.append(pid)
    ids.sort()
    return ids


def price_deltas(unit_base: float) -> tuple[float, float, float]:
    """Option value *price* column (added to product base with prefix +). Cart qty stays 1."""
    u = max(0.0, float(unit_base))
    deltas = []
    for n, pct in TIERS:
        factor = 1.0 - (pct / 100.0)
        line = n * u * factor
        delta = line - u
        deltas.append(round(delta, 4))
    return tuple(deltas)


def build_options_sheet() -> list[list]:
    row0 = ["option_id", "type", "sort_order"] + [f"name({c})" for c in LANGS]
    row1 = [OPTION_ID, "radio", 1] + [OPTION_NAMES[c] for c in LANGS]
    return [row0, row1]


def build_option_values_sheet() -> list[list]:
    row0 = ["option_value_id", "option_id", "image", "sort_order"] + [f"name({c})" for c in LANGS]
    rows = [row0]
    for idx, vid in enumerate(OPTION_VALUE_IDS):
        r = [vid, OPTION_ID, "", idx + 1] + [VALUE_NAMES[idx][c] for c in LANGS]
        rows.append(r)
    return rows


def build_product_options_rows(product_ids: list[int]) -> list[list]:
    default_val = VALUE_NAMES[0][DEFAULT_LANG_CODE]
    row0 = ["product_id", "option", "default_option_value", "required"]
    rows = [row0]
    for pid in product_ids:
        rows.append([pid, OPTION_NAMES[DEFAULT_LANG_CODE], default_val, "1"])
    return rows


def build_product_option_values_rows(ws_products, headers: list, product_ids: list[int]) -> list[list]:
    hi = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}
    pi, price_i = hi["product_id"], hi["price"]
    row0 = [
        "product_id",
        "option",
        "option_value",
        "quantity",
        "subtract",
        "price",
        "price_prefix",
        "points",
        "points_prefix",
        "weight",
        "weight_prefix",
    ]
    rows = [row0]
    price_by_pid: dict[int, float] = {}
    for row in ws_products.iter_rows(min_row=2, values_only=True):
        if not row or row[pi] is None:
            continue
        try:
            pid = int(row[pi])
        except (TypeError, ValueError):
            continue
        if pid not in product_ids:
            continue
        try:
            price_by_pid[pid] = float(row[price_i] or 0)
        except (TypeError, ValueError):
            price_by_pid[pid] = 0.0

    for pid in product_ids:
        unit = price_by_pid.get(pid, 0.0)
        d0, d1, d2 = price_deltas(unit)
        deltas = [d0, d1, d2]
        for idx, valdef in enumerate(VALUE_NAMES):
            rows.append(
                [
                    pid,
                    OPTION_NAMES[DEFAULT_LANG_CODE],
                    valdef[DEFAULT_LANG_CODE],
                    999,
                    "false",
                    deltas[idx],
                    "+",
                    0,
                    "+",
                    "0.00",
                    "+",
                ]
            )
    return rows


def write_sheet(wb: openpyxl.Workbook, name: str, data: list[list], index: int | None = None) -> None:
    if name in wb.sheetnames:
        wb.remove(wb[name])
    ws = wb.create_sheet(name, index if index is not None else len(wb.sheetnames))
    for r, row in enumerate(data, start=1):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)


def strip_volume_discount_rows(wb: openpyxl.Workbook) -> int:
    """
    Remove Discounts rows with quantity > 1 (legacy volume tiers). Keeps qty=1 rows (specials / sale lines).
    """
    if "Discounts" not in wb.sheetnames:
        return 0
    ws = wb["Discounts"]
    header = [c.value for c in ws[1]]
    idx_q = None
    for i, h in enumerate(header):
        if h is not None and str(h).strip().lower() == "quantity":
            idx_q = i
            break
    if idx_q is None:
        return 0
    col = idx_q + 1
    removed = 0
    for r in range(ws.max_row, 1, -1):
        raw = ws.cell(row=r, column=col).value
        try:
            q = float(raw) if raw is not None and str(raw).strip() != "" else 0.0
        except (TypeError, ValueError):
            q = 0.0
        if q > 1.0:
            ws.delete_rows(r)
            removed += 1
    return removed


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "workbook",
        nargs="?",
        default=str(Path(__file__).resolve().parents[1] / "products_import.xlsx"),
    )
    args = parser.parse_args()
    path = Path(args.workbook)
    if not path.is_file():
        raise SystemExit(f"Missing workbook: {path}")

    wb = openpyxl.load_workbook(path)
    if "Products" not in wb.sheetnames:
        raise SystemExit("Workbook has no Products sheet")

    ws_p = wb["Products"]
    headers = [c.value for c in ws_p[1]]
    product_ids = product_ids_for_categories(ws_p, headers)

    vol_removed = strip_volume_discount_rows(wb)

    # Insert global option sheets after Products (import order in extension is Options → OptionValues before ProductOptions)
    idx_after_products = wb.sheetnames.index("Products") + 1
    write_sheet(wb, "Options", build_options_sheet(), idx_after_products)
    idx_ov = wb.sheetnames.index("Options") + 1
    write_sheet(wb, "OptionValues", build_option_values_sheet(), idx_ov)
    idx_po = wb.sheetnames.index("OptionValues") + 1
    write_sheet(wb, "ProductOptions", build_product_options_rows(product_ids), idx_po)
    idx_pov = wb.sheetnames.index("ProductOptions") + 1
    write_sheet(
        wb,
        "ProductOptionValues",
        build_product_option_values_rows(ws_p, headers, product_ids),
        idx_pov,
    )

    wb.save(path)
    print(
        f"Updated {path.name}: Options + OptionValues + ProductOptions + ProductOptionValues "
        f"for {len(product_ids)} products; Discounts sheet preserved"
        + (f", removed {vol_removed} volume (qty>1) discount row(s)" if vol_removed else "")
        + "."
    )


if __name__ == "__main__":
    main()
